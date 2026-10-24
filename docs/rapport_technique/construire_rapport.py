"""Construit le rapport technique Word a partir du brouillon Markdown, dans le modele officiel.

    python docs/rapport_technique/construire_rapport.py [--corps 10.5] [--tableaux 8.5]

Le modele est le rapport Word existant (styles, en-tete, titres numerotes,
figure) ; son corps est remplace par le contenu de brouillon_v2_EN.md :
« ## N. Titre » -> titre numerote, paragraphes, listes « - », tableaux « | »,
**gras**. La figure 1 est remplacee par figure_architecture.png et inseree
apres la section 2.
"""

from __future__ import annotations

import argparse
import copy
import re
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
MODELE = HERE / "MinorShield_Policy_Lab_Technical_Report.docx"
BROUILLON = HERE / "brouillon_v2_EN.md"
FIGURE = HERE / "figure_architecture.png"
SORTIE = HERE / "MinorShield_Policy_Lab_Technical_Report.docx"


def runs_from_markdown(par, text: str, size: float) -> None:
    """Ajoute le texte au paragraphe en gerant **gras** ; les `code` deviennent du texte simple."""
    text = text.replace("`", "")
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            r = par.add_run(part)
            r.bold = i % 2 == 1
            r.font.size = Pt(size)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--corps", type=float, default=10.5)
    ap.add_argument("--tableaux", type=float, default=8.5)
    ap.add_argument("--modele", type=Path, default=MODELE)
    ap.add_argument("--sortie", type=Path, default=SORTIE)
    ap.add_argument("--brouillon", type=Path, default=BROUILLON)
    ap.add_argument("--figure", type=Path, default=FIGURE)
    ap.add_argument("--legende", default="Figure 1 – Architecture of MinorShield Policy Lab mapped onto the ITU-T "
                                         "Y.3172 pipeline, with the proposed verification node")
    args = ap.parse_args()

    d = docx.Document(str(args.modele))
    body = d.element.body
    kids = list(body.iterchildren())
    sect = kids[-1]
    heading_tpl = next(k for k in kids if k.find(f".//{qn('w:numPr')}") is not None)
    para_tpl = next(k for k in kids[4:] if k.tag == qn("w:p") and k.find(f".//{qn('w:numPr')}") is None
                    and k.find(f".//{qn('a:blip')}") is None and "".join(k.itertext()).strip())
    fig_tpl = next(k for k in kids if k.find(f".//{qn('a:blip')}") is not None)
    cap_tpl = kids[kids.index(fig_tpl) + 1]
    tbl_style = next((docx.table.Table(k, d).style for k in kids if k.tag == qn("w:tbl")), None)

    # Image : nouveau contenu, proportions recalculees.
    blip = fig_tpl.find(f".//{qn('a:blip')}")
    img_part = d.part.related_parts[blip.get(qn("r:embed"))]
    png = args.figure.read_bytes()
    img_part._blob = png
    from PIL import Image
    with Image.open(args.figure) as im:
        w, h = im.size
    for ext in fig_tpl.iter():
        if ext.tag in (qn("wp:extent"), qn("a:ext")) and ext.get("cx"):
            ext.set("cy", str(int(int(ext.get("cx")) * h / w)))

    # On garde l'en-tete (4 premiers paragraphes), on vide le reste.
    for k in kids[4:-1]:
        body.remove(k)

    def new_par(text: str, size: float):
        p_el = copy.deepcopy(para_tpl)
        for r in p_el.findall(qn("w:r")):
            p_el.remove(r)
        sect.addprevious(p_el)
        par = docx.text.paragraph.Paragraph(p_el, d)
        runs_from_markdown(par, text, size)
        return par

    def new_heading(text: str):
        h_el = copy.deepcopy(heading_tpl)
        for r in h_el.findall(qn("w:r"))[1:]:
            h_el.remove(r)
        h_el.findall(f".//{qn('w:t')}")[0].text = text
        sect.addprevious(h_el)

    def new_table(rows: list[list[str]]):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        if tbl_style is not None:
            t.style = tbl_style
        # Largeur des colonnes proportionnelle a la longueur moyenne du texte.
        sec = d.sections[0]
        usable = sec.page_width - sec.left_margin - sec.right_margin
        weights = [max(sum(len(r[j]) for r in rows) / len(rows), 10) for j in range(len(rows[0]))]
        widths = [int(usable * wgt / sum(weights)) for wgt in weights]
        t.autofit = False
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.width = widths[j]
                cell.text = ""
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                runs_from_markdown(par, f"**{val}**" if i == 0 and "**" not in val else val, args.tableaux)
        sect.addprevious(t._tbl)
        new_par("", 4)   # petit espace apres le tableau

    lines = args.brouillon.read_text(encoding="utf-8").splitlines()
    start = next(i for i, l in enumerate(lines) if l.startswith("## 1."))
    table: list[list[str]] = []
    section = 0
    for line in lines[start:] + [""]:
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                table.append(cells)
            continue
        if table:
            new_table(table)
            table = []
        if line.startswith("## "):
            if section == 2:   # figure apres la section 2 (cas d'usage)
                sect.addprevious(copy.deepcopy(fig_tpl))
                cap = copy.deepcopy(cap_tpl)
                sect.addprevious(cap)
            section += 1
            new_heading(re.sub(r"^##\s+\d+\.\s*", "", line))
        elif line.startswith("- "):
            new_par("• " + line[2:], args.corps)
        elif line.strip():
            new_par(line, args.corps)

    # Legende de la figure, a jour.
    for p in d.paragraphs:
        if p.text.startswith("Figure 1"):
            for r in p.runs[1:]:
                r.text = ""
            p.runs[0].text = args.legende
    d.save(str(args.sortie))
    print(args.sortie)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
